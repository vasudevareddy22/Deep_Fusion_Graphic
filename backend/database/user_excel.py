"""
User Excel Logger — DeepFusionGuard
Maintains registered_users.xlsx with styled headers.
All registration and login events are logged here.
"""

import os
import hashlib
import uuid
from datetime import datetime
from threading import Lock

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# Excel file location — same directory as this file
EXCEL_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(EXCEL_DIR, "registered_users.xlsx")

COLUMNS = [
    "User ID",
    "Full Name",
    "Identifier",          # email or mobile number
    "Auth Method",         # PASSWORD | EMAIL_OTP | MOBILE_OTP | GOOGLE
    "Organization",
    "Role",
    "Registration Timestamp",
    "Credential Hash",     # SHA-256 of password (empty for OTP/Google)
    "Status",              # ACTIVE | SUSPENDED
    "Last Login"
]

HEADER_BG = "1A237E"       # deep navy blue
HEADER_FONT_COLOR = "FFFFFF"

_excel_lock = Lock()


def _init_excel():
    """Create the Excel file with styled header row if it doesn't exist."""
    if not OPENPYXL_AVAILABLE:
        return
    if os.path.exists(EXCEL_PATH):
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registered Users"

    # Write header
    ws.append(COLUMNS)
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR, size=11)
        cell.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="FFFFFF")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Set column widths
    widths = [15, 22, 30, 16, 22, 14, 24, 66, 12, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24

    # Freeze header
    ws.freeze_panes = "A2"
    wb.save(EXCEL_PATH)


def _hash_credential(credential: str) -> str:
    """SHA-256 hash of credential string."""
    if not credential:
        return ""
    return hashlib.sha256(credential.encode("utf-8")).hexdigest()


def log_registration(
    name: str,
    identifier: str,
    auth_method: str,
    organization: str = "",
    role: str = "analyst",
    credential: str = ""
) -> str:
    """
    Add a new user row to the Excel file.
    Returns the generated User ID (UUID4 short).
    """
    if not OPENPYXL_AVAILABLE:
        return str(uuid.uuid4())[:8].upper()

    _init_excel()
    user_id = str(uuid.uuid4())[:8].upper()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cred_hash = _hash_credential(credential) if credential else ""

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
            ws.append([
                user_id,
                name,
                identifier,
                auth_method.upper(),
                organization,
                role,
                now_str,
                cred_hash,
                "ACTIVE",
                ""            # Last Login filled on first login
            ])
            wb.save(EXCEL_PATH)
        except Exception as e:
            print(f"[UserExcel] Registration log error: {e}")

    return user_id


def log_login(identifier: str) -> None:
    """Update Last Login timestamp for the given identifier."""
    if not OPENPYXL_AVAILABLE:
        return

    _init_excel()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    identifier = identifier.strip().lower()

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
            id_col = COLUMNS.index("Identifier") + 1
            login_col = COLUMNS.index("Last Login") + 1
            for row in ws.iter_rows(min_row=2):
                cell_val = row[id_col - 1].value
                if cell_val and str(cell_val).strip().lower() == identifier:
                    row[login_col - 1].value = now_str
                    break
            wb.save(EXCEL_PATH)
        except Exception as e:
            print(f"[UserExcel] Login log error: {e}")


def user_exists(identifier: str) -> bool:
    """Check if an identifier already exists in the Excel file."""
    if not OPENPYXL_AVAILABLE:
        return False

    _init_excel()
    identifier = identifier.strip().lower()

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
            ws = wb.active
            id_col = COLUMNS.index("Identifier") + 1
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[id_col - 1] and str(row[id_col - 1]).strip().lower() == identifier:
                    return True
        except Exception as e:
            print(f"[UserExcel] user_exists error: {e}")
    return False


def get_all_users() -> list:
    """Return all user rows as list of dicts (excluding credential hash)."""
    if not OPENPYXL_AVAILABLE:
        return []

    _init_excel()
    users = []
    safe_cols = [c for c in COLUMNS if c != "Credential Hash"]

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:
                    continue
                user = {}
                for i, col in enumerate(COLUMNS):
                    if col != "Credential Hash":
                        user[col] = row[i] if i < len(row) else ""
                users.append(user)
        except Exception as e:
            print(f"[UserExcel] get_all_users error: {e}")
    return users


def get_excel_path() -> str:
    """Return the absolute path to the Excel file."""
    _init_excel()
    return EXCEL_PATH


def delete_user_from_excel(identifier: str) -> bool:
    """Delete a user row from registered_users.xlsx by identifier (email or mobile)."""
    if not OPENPYXL_AVAILABLE:
        return False
    _init_excel()
    identifier = identifier.strip().lower()
    deleted = False

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
            id_col = COLUMNS.index("Identifier") + 1
            row_to_delete = None
            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                cell_val = row[id_col - 1].value
                if cell_val and str(cell_val).strip().lower() == identifier:
                    row_to_delete = idx
                    break
            if row_to_delete:
                ws.delete_rows(row_to_delete)
                wb.save(EXCEL_PATH)
                deleted = True
        except Exception as e:
            print(f"[UserExcel] Delete user error: {e}")
    return deleted


def delete_users_bulk_from_excel(identifiers: list) -> int:
    """Delete multiple user rows from registered_users.xlsx."""
    if not OPENPYXL_AVAILABLE:
        return 0
    _init_excel()
    targets = set(i.strip().lower() for i in identifiers if i)
    deleted_count = 0

    with _excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
            id_col = COLUMNS.index("Identifier") + 1
            
            # Iterate backwards so row deletions don't shift preceding target indices
            rows_to_delete = []
            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                cell_val = row[id_col - 1].value
                if cell_val and str(cell_val).strip().lower() in targets:
                    rows_to_delete.append(idx)
            
            for r_idx in reversed(rows_to_delete):
                ws.delete_rows(r_idx)
                deleted_count += 1
                
            if deleted_count > 0:
                wb.save(EXCEL_PATH)
        except Exception as e:
            print(f"[UserExcel] Bulk delete error: {e}")
    return deleted_count
